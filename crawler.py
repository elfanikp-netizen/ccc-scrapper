"""CCC workfile crawler driven by an input Excel workbook."""

from __future__ import annotations

import argparse
from collections import Counter, defaultdict, deque
import json
import logging
import os
import re
import sys
import tempfile
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from dotenv import load_dotenv
from openpyxl import Workbook, load_workbook
from playwright.sync_api import BrowserContext, Error as PlaywrightError, Locator, Page, TimeoutError as PlaywrightTimeoutError, sync_playwright


BASE_URL = "https://www.cccone.com/login/login-idp"
AUTH_STATE = Path("playwright/.auth/ccc-state.json")
OUTPUT_COLUMNS = [
    "VIN",
    "OEM",
    "Ipart listed on CCC",
    "Ipart List Price CCC",
    "Lowest CCC Price",
    "Lowest CCC Supplier",
    "Next Lowest CCC Price",
    "Next Lowest CCC Supplier",
    "Lowest CCC Price Supplier",
    "Status",
    "date_time",
    "market_data",
]
IPART_VENDOR = "iPart Automotive Inc."
TARGET_DESCRIPTION = "A/M: CAPA"
MAX_VENDOR_SEARCH_RETRIES = 3
RESULT_SAVE_INTERVAL = 1
RECALL_NOTICE_DETECT_TIMEOUT_MS = 1000
MODEL_MISMATCH_RE = re.compile(
    r"part does not match (?:(?:the )?(?:year[, ]+)?(?:model|engine)(?: of the vehicle)?|the year of the vehicle)",
    re.I,
)
NO_SOURCE_PARTS_RE = re.compile(
    r"no part lines to source|source parts\s*[:\-]?\s*0|0\s*source parts|enter part price",
    re.I,
)
MARKET_DATA_FIELDS = {
    "description": "description",
    "vendor": "vendor",
    "partno": "part-no",
    "partnumber": "part-no",
    "delivery": "delivery",
    "extprice": "ext-price",
    "labor": "labor",
    "paint": "paint",
    "total": "total",
    "diff": "diff",
    "margin": "margin",
}

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s", datefmt="%H:%M:%S")
log = logging.getLogger("ccc-crawler")


@dataclass
class PartOffer:
    vendor: str
    description: str
    total: float


class CatalogPartNotFoundError(RuntimeError):
    pass


class FatalCrawlerError(RuntimeError):
    pass


def block_nonessential_resources(route: Any) -> None:
    if route.request.resource_type in {"image", "media"}:
        route.abort()
    else:
        route.continue_()


def normalize(value: str) -> str:
    return re.sub(r"\s+", " ", value or "").strip().casefold()


def parse_price(value: str) -> float | None:
    cleaned = re.sub(r"[^0-9.\-]", "", value.replace(",", ""))
    try:
        return float(cleaned)
    except ValueError:
        return None


def read_input(path: Path) -> list[dict[str, str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    headers = [normalize(str(value or "")) for value in next(rows, ())]
    try:
        vin_index = headers.index("vin")
        oem_index = headers.index("oem")
    except ValueError as error:
        raise ValueError("Input workbook must have VIN and OEM column headers.") from error

    records = []
    for row_number, row in enumerate(rows, start=2):
        vin = str(row[vin_index] or "").strip() if len(row) > vin_index else ""
        oem_cell = str(row[oem_index] or "").strip() if len(row) > oem_index else ""
        oems = [value.strip() for value in oem_cell.split(",") if value.strip()]
        if vin and oems:
            for oem_index_in_cell, oem in enumerate(oems, start=1):
                row_label = str(row_number) if len(oems) == 1 else f"{row_number}.{oem_index_in_cell}"
                records.append({"VIN": vin, "OEM": oem, "_row": row_label})
        elif vin or oem_cell:
            log.warning("Skipping input row %d: VIN or OEM is blank", row_number)
    workbook.close()
    return records


def save_results(path: Path, results: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "CCC Results"
    sheet.append(OUTPUT_COLUMNS)
    for result in results:
        sheet.append([result.get(column) for column in OUTPUT_COLUMNS])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    temporary_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            dir=path.parent,
            prefix=f".{path.stem}.",
            suffix=path.suffix,
            delete=False,
        ) as temporary_file:
            temporary_path = Path(temporary_file.name)
        workbook.save(temporary_path)
        with temporary_path.open("rb+") as saved_file:
            os.fsync(saved_file.fileno())
        os.replace(temporary_path, path)
    finally:
        workbook.close()
        if temporary_path is not None:
            temporary_path.unlink(missing_ok=True)


def load_results(path: Path) -> list[dict[str, Any]]:
    if not path.is_file():
        return []
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        headers = [str(value or "").strip() for value in next(rows, ())]
        header_indexes = {header: index for index, header in enumerate(headers)}
        if "VIN" not in header_indexes or "OEM" not in header_indexes:
            log.warning("Existing output has no VIN/OEM columns; starting without prior results")
            return []
        results = []
        for row in rows:
            result = {
                column: row[index] if index < len(row) else None
                for column, index in header_indexes.items()
                if column in OUTPUT_COLUMNS
            }
            if result.get("VIN") and result.get("OEM"):
                results.append(result)
        return results
    finally:
        workbook.close()


def resume_state(results: list[dict[str, Any]]) -> tuple[Counter, dict[tuple[str, str], deque[int]]]:
    completed: Counter = Counter()
    retry_rows: dict[tuple[str, str], deque[int]] = defaultdict(deque)
    for index, result in enumerate(results):
        key = (str(result.get("VIN") or "").strip(), str(result.get("OEM") or "").strip())
        if not all(key):
            continue
        status = str(result.get("Status") or "").strip()
        if status.upper().startswith("ERROR:"):
            retry_rows[key].append(index)
        else:
            completed[key] += 1
    return completed, retry_rows


class CCCCrawler:
    def __init__(
        self,
        page: Page,
        username: str,
        password: str,
        artifacts: Path,
        description_target: str = TARGET_DESCRIPTION,
        context: BrowserContext | None = None,
    ):
        self.page = page
        self.context = context or page.context
        self.username = username
        self.password = password
        self.artifacts = artifacts
        self.description_target = description_target
        self.workfiles_url: str | None = None
        self._workfile_initialized = False
        self.artifacts.mkdir(parents=True, exist_ok=True)

    def checkpoint(self, label: str) -> None:
        log.info("%s", label)

    @staticmethod
    def first_visible(locator: Locator) -> Locator | None:
        for index in range(locator.count()):
            candidate = locator.nth(index)
            try:
                if candidate.is_visible():
                    return candidate
            except Exception:
                continue
        return None

    def find_textbox(self, hints: tuple[str, ...]) -> Locator:
        for hint in hints:
            candidates = [
                self.page.get_by_label(re.compile(hint, re.I)),
                self.page.get_by_placeholder(re.compile(hint, re.I)),
                self.page.locator(f'input[aria-label*="{hint}" i]'),
                self.page.locator(f'input[name*="{hint}" i]'),
            ]
            for candidate in candidates:
                visible = self.first_visible(candidate)
                if visible is not None:
                    return visible
        visible_inputs = self.page.locator('input:visible:not([type="hidden"]):not([type="checkbox"])')
        if visible_inputs.count() == 1:
            return visible_inputs.first
        raise RuntimeError(f"Could not identify a visible text box using hints: {hints}")

    def click_text(self, label: str, *, exact: bool = True, timeout: int = 8000) -> None:
        candidates = [
            self.page.get_by_role("button", name=re.compile(re.escape(label), re.I)),
            self.page.get_by_role("link", name=re.compile(re.escape(label), re.I)),
            self.page.get_by_text(label, exact=exact),
        ]
        for candidate in candidates:
            visible = self.first_visible(candidate)
            if visible is not None:
                try:
                    visible.click(timeout=timeout)
                    return
                except Exception:
                    continue
        raise RuntimeError(f"Could not find clickable control: {label}")

    def wait_for_textbox(self, hints: tuple[str, ...], timeout: int = 30000) -> Locator:
        try:
            self.page.wait_for_function(
                """hints => [...document.querySelectorAll('input')].some(input => {
                    if (!input.offsetWidth && !input.offsetHeight) return false;
                    const text = [input.placeholder, input.getAttribute('aria-label'), input.name, input.id].join(' ');
                    return hints.some(hint => new RegExp(hint, 'i').test(text));
                })""",
                list(hints),
                timeout=timeout,
            )
        except PlaywrightTimeoutError:
            pass
        return self.find_textbox(hints)

    def capture_failure(self, label: str) -> None:
        safe_label = re.sub(r"[^A-Za-z0-9_-]+", "_", label)[:60]
        try:
            self.page.screenshot(path=self.artifacts / f"{safe_label}.png", full_page=True)
            (self.artifacts / f"{safe_label}.txt").write_text(self.page.locator("body").inner_text()[:20000], encoding="utf-8")
            log.error("Saved failure evidence under %s", self.artifacts)
        except Exception as error:
            log.error("Could not save failure evidence: %s", error)

    def sign_in_if_needed(self) -> None:
        self.checkpoint("Opening CCC and checking sign-in state")
        self.page.goto(BASE_URL, wait_until="domcontentloaded")
        login_button = self.page.get_by_role("button", name=re.compile("^Login$", re.I))
        try:
            login_button.wait_for(state="visible", timeout=15000)
            self.checkpoint("CCCOne login landing page found; clicking Login")
            login_button.click()
        except PlaywrightTimeoutError:
            log.info("CCCOne login landing button not shown; checking the current session page")

        self.checkpoint(f"Identity page reached: {self.page.url}")
        self.checkpoint("Waiting for the sign-in form or authenticated navigation to load")
        try:
            self.page.wait_for_function(
                r"""() => {
                    const visible = element => !!element && !!(element.offsetWidth || element.offsetHeight);
                    const username = visible(document.querySelector('#idp-discovery-username, input[name="username"]'));
                    const password = [...document.querySelectorAll('input[type="password"]')].some(visible);
                    const body = (document.body?.innerText || '').toLowerCase();
                    const authenticatedNavigation = (location.hostname === 'www.cccone.com' && !location.pathname.startsWith('/login/'))
                        || (body.includes('workfiles') && !body.includes('sign in'));
                    return username || password || authenticatedNavigation;
                }""",
                timeout=30000,
            )
        except PlaywrightTimeoutError as error:
            self.capture_failure("ccc_entry_surface_not_loaded")
            raise RuntimeError("CCC did not show a login form or authenticated navigation within 30 seconds") from error

        password_box = self.page.locator('input[type="password"]:visible').first
        if not password_box.count():
            username_box = self.page.locator('#idp-discovery-username:visible, input[name="username"]:visible').first
            if not username_box.count():
                if self.page.url.startswith("https://www.cccone.com/") and "/login/" not in self.page.url:
                    log.info("Authenticated CCCOne application route detected; skipping sign-in")
                    return
                raise RuntimeError("CCC loaded, but neither a username field nor authenticated application route was found")

            username_box.fill(self.username)
            for attempt in range(5):
                if self.page.locator('input[type="password"]:visible').count():
                    break
                self.checkpoint(f"Advancing username sign-in step ({attempt + 1}/5)")
                self.click_text("Next")
                try:
                    self.page.locator('input[type="password"]').wait_for(state="visible", timeout=8000)
                except PlaywrightTimeoutError:
                    continue
            password_box = self.page.locator('input[type="password"]:visible').first
        if not password_box.count():
            raise RuntimeError("Password field did not appear after advancing the sign-in flow")
        if password_box.count():
            password_box.fill(self.password)
            self.click_text("Sign in")
            self.page.wait_for_load_state("domcontentloaded")
            try:
                self.page.locator('input[type="password"]:visible').wait_for(state="hidden", timeout=30000)
            except PlaywrightTimeoutError:
                log.warning("Password field remains visible; sign-in may require an additional prompt")
        log.info("Sign-in flow completed or an existing session was detected")

    def open_workfiles(self) -> None:
        self.checkpoint("Opening CCCOne navigation menu")
        self.page.get_by_role("button", name="Nav Menu").click()
        workfiles = self.page.get_by_text("Workfiles", exact=True)
        workfiles.first.wait_for(state="visible", timeout=15000)
        workfiles.first.click()
        self.page.wait_for_function(
            """() => [...document.querySelectorAll('*')].filter(element => {
                const rect = element.getBoundingClientRect();
                return rect.width && rect.height && element.children.length === 0 && element.textContent.trim() === 'Workfiles';
            }).length >= 2""",
            timeout=10000,
        )
        self.page.get_by_text("Workfiles", exact=True).last.click()
        self.page.wait_for_url("**/workfile-search/workfile", timeout=30000)
        self.workfiles_url = self.page.url
        self.checkpoint(f"Workfiles list opened: {self.workfiles_url}")

    def open_new_workfile(self) -> None:
        if not self.workfiles_url:
            raise RuntimeError("Workfiles list URL is not initialized")
        if self.page.url != self.workfiles_url:
            self.page.goto(self.workfiles_url, wait_until="domcontentloaded")
        self.dismiss_onboarding_guide()
        self.checkpoint("Creating a new workfile for this input row")
        self.page.get_by_text("New workfile", exact=True).click(timeout=15000, no_wait_after=True)
        try:
            self.page.wait_for_url("**/edit/contacts", wait_until="commit", timeout=45000)
            self.page.locator('a[aria-label="Vehicle"]').wait_for(state="visible", timeout=45000)
        except PlaywrightTimeoutError as error:
            self.capture_failure("new_workfile_surface_not_loaded")
            raise RuntimeError(
                f"CCC did not finish opening the workfile Vehicle panel; current URL: {self.page.url}"
            ) from error
        self.dismiss_jumpstart_prompt()
        self._workfile_initialized = True
        log.info("New workfile opened: %s", self.page.url)

    def dismiss_jumpstart_prompt(self) -> None:
        prompt = self.page.locator(".cdk-overlay-pane").filter(
            has_text=re.compile(r"What would you like to do first\?", re.I)
        ).last
        if not prompt.count():
            return
        jumpstart = prompt.get_by_text("Jumpstart estimate", exact=True)
        control = self.first_visible(jumpstart)
        if control is None:
            raise RuntimeError("CCC first-estimate prompt appeared without a Jumpstart estimate control")
        self.checkpoint("CCC first-estimate prompt shown; choosing Jumpstart estimate")
        control.click(timeout=10000)
        prompt.wait_for(state="hidden", timeout=15000)

    def dismiss_onboarding_guide(self) -> None:
        guide = self.page.locator("#pendo-base")
        if not guide.count():
            return
        maybe_later = guide.get_by_role("button", name=re.compile(r"^MAYBE LATER$", re.I))
        control = self.first_visible(maybe_later)
        if control is not None:
            control.click(timeout=2000)

    def return_to_workfiles(self) -> None:
        if self.workfiles_url:
            self.page.goto(self.workfiles_url, wait_until="domcontentloaded")

    def enter_vehicle(self, vin: str) -> bool:
        self.checkpoint(f"Vehicle: searching VIN {vin}")
        self.dismiss_jumpstart_prompt()
        self.page.locator('a[aria-label="Vehicle"]').click()
        self.open_vehicle_selector()
        vin_box = self.page.locator(".vin-input input.input-box:visible").last
        vin_box.wait_for(state="visible", timeout=10000)
        vin_box.fill(vin)
        self.page.get_by_text("Valid VIN", exact=True).wait_for(state="visible", timeout=20000)
        self.checkpoint("CCC validated the VIN; advancing to model selection")
        self.click_text("Next")
        selection_state = self.resolve_vehicle_change_state(self.page.url)
        while selection_state == "options":
            options = self.page.locator('[role="option"]:visible')
            options.first.click()
            self.checkpoint("CCC vehicle model selected; advancing to Finish")
            self.click_text("Next")
            selection_state = self.resolve_vehicle_change_state(self.page.url)
        if selection_state != "done":
            raise FatalCrawlerError(f"Unexpected CCC vehicle selection state: {selection_state}")
        self.checkpoint("CCC vehicle selection complete; clicking DONE")
        self.finish_vehicle_selection()
        log.info("Vehicle selected for VIN %s", vin)
        return True

    def resolve_vehicle_change_state(self, workfile_url: str) -> str:
        state = self.wait_for_vehicle_selection_state()
        while state == "change-confirmation":
            self.confirm_change_vehicle()
            try:
                state = self.wait_for_vehicle_selection_state()
            except PlaywrightError as error:
                if not self.page.is_closed() and "closed" not in str(error).lower():
                    raise
                self.reopen_workfile_page(workfile_url)
                state = self.wait_for_vehicle_selection_state()
        return state

    def wait_for_vehicle_selection_state(self) -> str:
        return self.page.wait_for_function(
            """() => {
                const visible = element => {
                    const rect = element.getBoundingClientRect();
                    const style = getComputedStyle(element);
                    return rect.width > 0 && rect.height > 0 && style.visibility !== 'hidden' && style.display !== 'none';
                };
                const changeConfirmation = [...document.querySelectorAll('.cdk-overlay-pane')].some(pane =>
                    visible(pane) && /Any estimate lines written will be erased[\\s\\S]*Change vehicle anyway\\?/i.test(pane.innerText || '')
                );
                if (changeConfirmation) return 'change-confirmation';
                const options = [...document.querySelectorAll('[role="option"]')].some(option =>
                    visible(option) && (option.innerText || option.textContent || '').trim().length > 0
                );
                if (options) return 'options';
                const finish = [...document.querySelectorAll('button')].some(button =>
                    visible(button) && /^(?:DONE|FINISH)$/i.test((button.innerText || button.textContent || '').trim()) && !button.disabled
                );
                return finish ? 'done' : false;
            }""",
            timeout=60000,
        ).json_value()

    def confirm_change_vehicle(self) -> None:
        try:
            self.page.get_by_text(
                re.compile(r"Any estimate lines written will be erased.*Change vehicle anyway\?", re.I | re.S)
            ).wait_for(state="visible", timeout=15000)
            prompt = self.page.locator(".cdk-overlay-pane").filter(
                has_text=re.compile(r"Any estimate lines written will be erased.*Change vehicle anyway\?", re.I | re.S)
            ).last
            prompt.wait_for(state="visible", timeout=10000)
            buttons = prompt.locator("button")
            change_button = buttons.filter(has_text=re.compile(r"^\s*CHANGE VEHICLE\s*$", re.I)).last
            change_button.wait_for(state="visible", timeout=10000)
            self.checkpoint("CCC warned that changing vehicle erases estimate lines; clicking CHANGE VEHICLE as requested")
            change_button.click(timeout=15000)
            self.checkpoint("CCC CHANGE VEHICLE confirmation clicked")
        except PlaywrightError as error:
            raise FatalCrawlerError(f"CCC Change Vehicle confirmation could not complete: {error}") from error

    def reopen_workfile_page(self, workfile_url: str) -> None:
        self.checkpoint("CCC closed the workfile tab during vehicle change; reopening the same workfile")
        pages = [page for page in self.context.pages if not page.is_closed()]
        self.page = pages[-1] if pages else self.context.new_page()
        if self.page.url != workfile_url:
            self.page.goto(workfile_url, wait_until="domcontentloaded", timeout=60000)
        try:
            self.page.locator('a[aria-label="Vehicle"]').wait_for(state="visible", timeout=45000)
        except PlaywrightTimeoutError as error:
            raise RuntimeError("CCC did not restore the existing workfile after vehicle change") from error

    def open_vehicle_selector(self) -> None:
        self.page.wait_for_function(
            """() => [...document.querySelectorAll('button')].some(button => {
                const rect = button.getBoundingClientRect();
                const style = getComputedStyle(button);
                const visible = rect.width > 0 && rect.height > 0 && style.visibility !== 'hidden' && style.display !== 'none';
                const label = (button.innerText || button.textContent || button.getAttribute('aria-label') || '').trim().toLowerCase();
                return visible && (label === 'select vehicle' || label === 'change vehicle');
            })""",
            timeout=30000,
        )
        buttons = self.page.locator("button")
        for index in range(buttons.count()):
            button = buttons.nth(index)
            label = normalize(button.inner_text() or button.get_attribute("aria-label") or "")
            if button.is_visible() and label in {"select vehicle", "change vehicle"}:
                self.checkpoint(f"Opening CCC vehicle selector with {label.title()}")
                button.click(timeout=15000)
                return
        raise RuntimeError("CCC Vehicle panel showed neither Select vehicle nor Change Vehicle")

    def finish_vehicle_selection(self) -> None:
        self.page.wait_for_function(
            """() => [...document.querySelectorAll('button')].some(button => {
                const rect = button.getBoundingClientRect();
                const style = getComputedStyle(button);
                const visible = rect.width > 0 && rect.height > 0 && style.visibility !== 'hidden' && style.display !== 'none';
                const label = (button.innerText || button.textContent || '').trim();
                return visible && /^(?:DONE|FINISH)$/i.test(label) && !button.disabled;
            })""",
            timeout=60000,
        )
        button_index = self.page.locator("button").evaluate_all(
            """buttons => buttons.findIndex(button => {
                const rect = button.getBoundingClientRect();
                const style = getComputedStyle(button);
                const visible = rect.width > 0 && rect.height > 0 && style.visibility !== 'hidden' && style.display !== 'none';
                const label = (button.innerText || button.textContent || '').trim();
                return visible && /^(?:DONE|FINISH)$/i.test(label) && !button.disabled;
            })"""
        )
        if button_index < 0:
            raise RuntimeError("CCC's visible DONE/FINISH button disappeared before it could be clicked")
        self.page.locator("button").nth(button_index).click(timeout=15000)
        self.page.wait_for_function(
            """() => ![...document.querySelectorAll('button')].some(button => {
                const rect = button.getBoundingClientRect();
                const style = getComputedStyle(button);
                const visible = rect.width > 0 && rect.height > 0 && style.visibility !== 'hidden' && style.display !== 'none';
                const label = (button.innerText || button.textContent || '').trim();
                return visible && /^(?:DONE|FINISH)$/i.test(label);
            })""",
            timeout=30000,
        )
        self.handle_vehicle_recall_prompts()

    def handle_vehicle_recall_prompts(self) -> None:
        recall_notice = self.page.get_by_text(re.compile("vehicle recall notice", re.I)).first
        try:
            recall_notice.wait_for(state="visible", timeout=RECALL_NOTICE_DETECT_TIMEOUT_MS)
        except PlaywrightTimeoutError:
            return

        self.checkpoint("Vehicle recall notice shown; acknowledging with OK")
        self.click_text("OK")

        print_notice = self.page.get_by_text(re.compile("print recall notice", re.I)).first
        try:
            print_notice.wait_for(state="visible", timeout=10000)
        except PlaywrightTimeoutError:
            log.info("No print-recall prompt appeared after acknowledging the recall notice")
            return

        self.checkpoint("Print recall notice prompt shown; choosing No")
        no_pattern = re.compile(r"^\s*NO\s*$", re.I)
        no_control = self.page.locator(".cdk-overlay-pane .dialog-actions button").filter(has_text=no_pattern)
        visible_control = self.first_visible(no_control)
        if visible_control is None:
            no_control = self.page.locator(".cdk-overlay-pane .dialog-actions").get_by_text(no_pattern)
            visible_control = self.first_visible(no_control)
        if visible_control is None:
            self.capture_failure("print_recall_no_control_not_found")
            raise RuntimeError("Print recall prompt appeared, but its No control was not found")
        visible_control.click()
        self.page.get_by_text(re.compile("Print recall notice", re.I)).wait_for(state="hidden", timeout=10000)

    def open_part_results(self, oem: str) -> int | None:
        self.checkpoint(f"Estimate: searching OEM {oem}")
        self.dismiss_jumpstart_prompt()
        self.page.locator('a[aria-label="Estimate"]').click()
        self.dismiss_jumpstart_prompt()
        filter_button = self.page.locator("button.filter-button:visible").first
        filter_button.wait_for(state="visible", timeout=20000)
        filter_button.click()
        keyword_box = self.page.locator("input.search-input:visible").last
        keyword_box.wait_for(state="visible", timeout=10000)
        keyword_box.fill(oem)
        self.click_text("ALL MODELS")
        self.dismiss_jumpstart_prompt()
        filter_dialog = self.page.locator(".cdk-overlay-pane").filter(
            has_text=re.compile(r"Filter parts database", re.I)
        ).last
        try:
            filter_dialog.wait_for(state="visible", timeout=5000)
            ok_button = filter_dialog.locator("button").filter(
                has_text=re.compile(r"^\s*OK\s*$", re.I)
            )
            ok_control = self.first_visible(ok_button)
            if ok_control is None:
                ok_control = self.first_visible(
                    filter_dialog.get_by_text(re.compile(r"^\s*OK\s*$", re.I))
                )
            if ok_control is None:
                raise RuntimeError("Visible exact OK button was not found")
            ok_control.click(timeout=10000)
            filter_dialog.wait_for(state="hidden", timeout=10000)
        except (PlaywrightError, RuntimeError) as error:
            cancel = self.first_visible(
                filter_dialog.get_by_text(re.compile(r"^\s*CANCEL\s*$", re.I))
            )
            if cancel is None:
                cancel = self.first_visible(self.page.get_by_text("CANCEL", exact=True))
            if cancel is not None:
                cancel.click(timeout=5000)
            raise RuntimeError(f"CCC OEM filter could not confirm {oem} with its OK button: {error}") from error
        categories = self.page.locator("button.group-list-item:visible")
        try:
            categories.first.wait_for(state="visible", timeout=5000)
        except PlaywrightTimeoutError as error:
            category_labels = self.page.locator(".part-group-name:visible, .group-list-item:visible, [class*=group-list-item]:visible")
            if not category_labels.count():
                raise CatalogPartNotFoundError(f"No catalog results found for OEM {oem}") from error

        parts = self.page.locator("app-part-selector-part-row:visible")
        category_names = categories.all_inner_texts() if categories.count() else []
        if not category_names:
            category_names = [
                text.strip()
                for text in self.page.locator("body").inner_text().splitlines()
                if text.strip() and text.strip().isupper() and len(text.strip()) < 80
            ]
        selected_parts = False
        for category_name in category_names:
            visible_categories = self.page.locator("button.group-list-item:visible")
            matching_category = None
            for index in range(visible_categories.count()):
                candidate = visible_categories.nth(index)
                if normalize(candidate.inner_text()) == normalize(category_name):
                    matching_category = candidate
                    break
            self.checkpoint(f"Checking catalog category {category_name.strip()}")
            if matching_category is not None:
                matching_category.click(force=True)
            else:
                category_text = self.page.get_by_text(category_name.strip(), exact=True)
                visible_text = [
                    category_text.nth(index)
                    for index in range(category_text.count())
                    if category_text.nth(index).is_visible()
                ]
                if not visible_text:
                    continue
                visible_text[-1].click(force=True)
            try:
                parts.first.wait_for(state="visible", timeout=5000)
                selected_parts = True
            except PlaywrightTimeoutError:
                pass
            for _ in range(10):
                if selected_parts or parts.count():
                    selected_parts = True
                    break
                subgroups = self.page.locator("button.subgroup-button:visible")
                collapsed_subgroups = [
                    subgroups.nth(index)
                    for index in range(subgroups.count())
                    if "expanded" not in (subgroups.nth(index).get_attribute("class") or "")
                ]
                if not collapsed_subgroups:
                    break
                self.checkpoint("Expanding a catalog subgroup")
                collapsed_subgroups[0].click()
                try:
                    parts.first.wait_for(state="visible", timeout=5000)
                    selected_parts = True
                    break
                except PlaywrightTimeoutError:
                    continue
            if selected_parts:
                break
            no_results = self.page.get_by_text(re.compile(r"No results found for search text:", re.I))
            if self.first_visible(no_results) is not None:
                raise CatalogPartNotFoundError(f"No catalog results found for OEM {oem}")
            selected_group = self.page.locator("button.selected-group-button:visible")
            if selected_group.count():
                selected_group.click(force=True)
                try:
                    state = self.page.wait_for_function(
                        """() => {
                            const visible = element => !!element && !!(element.offsetWidth || element.offsetHeight);
                            if ([...document.querySelectorAll('button.group-list-item')].some(visible)) return 'categories';
                            if (/No results found for search text:/i.test(document.body?.innerText || '')) return 'empty';
                            return false;
                        }""",
                        timeout=10000,
                    ).json_value()
                except PlaywrightTimeoutError:
                    state = None
                if state == "empty":
                    raise CatalogPartNotFoundError(f"No catalog results found for OEM {oem}")
        if not selected_parts or not parts.count():
            raise CatalogPartNotFoundError(f"No catalog results found for OEM {oem} in any returned category")

        self.checkpoint("Double-clicking the first filtered catalog part")
        parts.first.dblclick()
        self.resolve_part_prompts()
        body_text = normalize(self.page.locator("body").inner_text())
        no_source_parts = bool(NO_SOURCE_PARTS_RE.search(body_text))
        source_tab = self.page.locator("[role=tab].source-parts-tab")
        if not source_tab.count() and no_source_parts:
            log.warning("CCC created the estimate line for OEM %s without sourceable parts; skipping vendor lookup", oem)
            return None
        try:
            source_tab.wait_for(state="visible", timeout=10000)
        except PlaywrightTimeoutError:
            if no_source_parts:
                log.warning("CCC created the estimate line for OEM %s without sourceable parts; skipping vendor lookup", oem)
                return None
            raise RuntimeError(f"Part selection did not complete for OEM {oem}; Source parts tab never appeared") from None
        try:
            self.page.wait_for_function(
                """() => {
                    const tab = document.querySelector('[role="tab"].source-parts-tab');
                    if (!tab) return 'missing';
                    return tab.getAttribute('aria-disabled') === 'false' ? 'enabled' : 'disabled';
                }""",
                timeout=15000,
            ).json_value()
        except PlaywrightTimeoutError:
            if no_source_parts:
                log.warning("CCC created the estimate line for OEM %s without sourceable parts; skipping vendor lookup", oem)
                return None
            log.warning("CCC kept Source parts disabled for OEM %s; skipping vendor lookup", oem)
            raise RuntimeError(f"CCC Source parts stayed disabled for OEM {oem}; retry this row") from None
        source_tab.click()
        source_tag = self.page.get_by_role("button", name="Show sourced parts").first
        try:
            source_tag.wait_for(state="visible", timeout=10000)
        except PlaywrightTimeoutError as error:
            body_text = normalize(self.page.locator("body").inner_text())
            if NO_SOURCE_PARTS_RE.search(body_text):
                log.info("CCC estimate line for OEM %s has no sourceable parts; skipping vendor lookup", oem)
                return None
            raise RuntimeError(f"CCC added no sourceable estimate line for OEM {oem}") from error
        source_tag.click()
        return self.wait_for_vendor_inventory()

    def wait_for_vendor_inventory(self) -> int | None:
        retry_count = 0
        self.vendor_search_retry_count = retry_count
        retry_button = self.page.get_by_role("button", name=re.compile(r"retry|try again", re.I))

        while True:
            self.checkpoint("Waiting for CCC vendor inventory search")
            try:
                outcome = self.page.wait_for_function(
                    """() => {
                        const visible = element => !!element && !!(element.offsetWidth || element.offsetHeight);
                        if ([...document.querySelectorAll('.cdk-overlay-pane .table-row')].some(visible)) return 'offers';
                        if ([...document.querySelectorAll('button')].some(button => visible(button) && /retry|try again/i.test(button.innerText))) return 'retry';
                        return false;
                    }""",
                    timeout=60000,
                ).json_value()
            except PlaywrightTimeoutError:
                log.error("CCC vendor search produced no offers or Retry control within 60 seconds")
                return None

            if outcome == "offers":
                return retry_count
            retry_control = self.first_visible(retry_button)
            if retry_control is None:
                log.error("CCC exposed a Retry outcome but its button was not accessible")
                return None
            if retry_count >= MAX_VENDOR_SEARCH_RETRIES:
                log.error("CCC vendor search still has no offers after %d retries", retry_count)
                return None

            retry_count += 1
            self.vendor_search_retry_count = retry_count
            self.checkpoint(f"Retrying CCC vendor inventory search ({retry_count}/{MAX_VENDOR_SEARCH_RETRIES})")
            retry_control.click()

    def resolve_part_prompts(self) -> None:
        for _ in range(4):
            try:
                self.page.wait_for_function(
                    """() => {
                        const text = (document.body?.innerText || '').toLowerCase();
                        return (text.includes('part does not match') && (text.includes('model') || text.includes('year')))
                            || text.includes('refinish this part')
                            || text.includes('select additional labor')
                            || text.includes('removing duplicate part');
                    }""",
                    timeout=3000,
                )
            except PlaywrightTimeoutError:
                return
            prompt = normalize(self.page.locator("body").inner_text())
            if MODEL_MISMATCH_RE.search(prompt):
                self.page.get_by_role("button", name=re.compile("^Yes$", re.I)).last.click()
                self.checkpoint("Accepted CCC model-mismatch confirmation for the selected part")
                self.page.get_by_text(MODEL_MISMATCH_RE).wait_for(state="hidden", timeout=10000)
            elif "refinish this part" in prompt:
                self.page.get_by_role("button", name=re.compile("^No$", re.I)).last.click()
                self.page.get_by_text(re.compile("Refinish this part", re.I)).wait_for(state="hidden", timeout=10000)
            elif "select additional labor" in prompt:
                self.page.get_by_role("button", name=re.compile("^OK$", re.I)).last.click()
                self.page.get_by_text(re.compile("Select additional labor", re.I)).wait_for(state="hidden", timeout=10000)
            elif "removing duplicate part" in prompt:
                self.page.get_by_role("button", name=re.compile("^OK$", re.I)).last.click()
                self.checkpoint("Acknowledged CCC duplicate-part notice")
                self.page.get_by_text(re.compile("Removing duplicate part", re.I)).wait_for(state="hidden", timeout=10000)
            else:
                return

    def read_vendor_inventory(self) -> tuple[list[PartOffer], list[dict[str, str]]]:
        rows = self.page.locator(".cdk-overlay-pane .table-row").evaluate_all(
            """rows => {
                const fields = {
                    description: 'description', vendor: 'vendor', partno: 'part-no',
                    partnumber: 'part-no', delivery: 'delivery', extprice: 'ext-price',
                    labor: 'labor', paint: 'paint', total: 'total', diff: 'diff', margin: 'margin'
                };
                const outputFields = [...new Set(Object.values(fields))];
                const attributes = ['data-full-value', 'data-original-value', 'data-value', 'title', 'aria-label'];
                return rows.map(row => {
                    const data = Object.fromEntries(outputFields.map(field => [field, '']));
                    for (const cell of row.querySelectorAll('[class*="column-"]')) {
                        const columnClass = [...cell.classList].find(name => name.startsWith('column-')) || '';
                        const key = columnClass.slice(7).toLowerCase().replace(/[^a-z0-9]/g, '');
                        const field = fields[key];
                        if (!field) continue;
                        let value = (cell.innerText || cell.textContent || '').trim();
                        if (field === 'part-no') {
                            const candidates = [cell, ...cell.querySelectorAll('*')].flatMap(node =>
                                attributes.map(name => node.getAttribute(name)).filter(candidate => candidate && candidate.trim())
                            ).map(candidate => candidate.trim()).filter(candidate =>
                                candidate.length > value.length && !/(?:\\.\\.\\.|…)/.test(candidate)
                            );
                            if (candidates.length) value = candidates.reduce((best, candidate) => candidate.length > best.length ? candidate : best, value);
                        }
                        data[field] = value;
                    }
                    return data;
                }).filter(data => Object.values(data).some(Boolean));
            }"""
        )
        market_data = sorted(rows, key=lambda item: normalize(item.get("description", "")))
        offers = []
        for row in market_data:
            total = parse_price(row["ext-price"])
            if row["description"] and row["vendor"] and total is not None:
                offers.append(PartOffer(row["vendor"], row["description"], total))
        if not offers:
            raise RuntimeError("CCC vendor result grid contained no readable Vendor/Description/EXT. PRICE rows")
        self.close_vendor_inventory()
        return offers, market_data

    def close_vendor_inventory(self) -> None:
        self.checkpoint("Closing CCC vendor inventory with Done")
        done = self.page.locator(".cdk-overlay-pane button").filter(
            has_text=re.compile(r"^\s*DONE\s*$", re.I)
        )
        try:
            done.first.wait_for(state="visible", timeout=10000)
            done.first.click(timeout=10000)
        except PlaywrightTimeoutError:
            raise RuntimeError("CCC vendor inventory results loaded, but the modal Done button was not found") from None
        self.page.locator(".cdk-overlay-pane").wait_for(state="hidden", timeout=15000)
        self.checkpoint("CCC vendor inventory modal closed")

    @staticmethod
    def summarize(
        vin: str,
        oem: str,
        offers: list[PartOffer],
        description_target: str = TARGET_DESCRIPTION,
    ) -> dict[str, Any]:
        matching_offers = [offer for offer in offers if normalize(description_target) in normalize(offer.description)]
        ipart = [offer for offer in matching_offers if normalize(offer.vendor) == normalize(IPART_VENDOR)]
        ipart_available = bool(ipart)
        sorted_all = sorted(matching_offers, key=lambda offer: offer.total)
        lowest_overall = sorted_all[0] if sorted_all else None
        second_lowest = sorted_all[1] if len(sorted_all) > 1 else None
        ipart_lowest = min(ipart, key=lambda offer: offer.total) if ipart else None
        return {
            "VIN": vin,
            "OEM": oem,
            "Ipart listed on CCC": ipart_available,
            "Ipart List Price CCC": ipart_lowest.total if ipart_lowest else None,
            "Lowest CCC Price": lowest_overall.total if lowest_overall else None,
            "Lowest CCC Supplier": lowest_overall.vendor if lowest_overall else None,
            "Next Lowest CCC Price": second_lowest.total if second_lowest else None,
            "Next Lowest CCC Supplier": second_lowest.vendor if second_lowest else None,
            "Lowest CCC Price Supplier": lowest_overall.vendor if lowest_overall else None,
            "Status": "OK" if matching_offers else f"No offers matching description: {description_target}",
        }

    def crawl_row(self, record: dict[str, str]) -> dict[str, Any]:
        vin, oem = record["VIN"], record["OEM"]
        self.checkpoint(f"Input row {record['_row']}: VIN={vin}, OEM={oem}")
        result: dict[str, Any] = {"VIN": vin, "OEM": oem}
        try:
            if not self._workfile_initialized:
                self.open_new_workfile()
            elif not self.page.locator('a[aria-label="Vehicle"]').count():
                if self.workfiles_url:
                    self.page.goto(self.workfiles_url, wait_until="domcontentloaded")
                self.open_new_workfile()
            if not self.enter_vehicle(vin):
                result["Status"] = "Vehicle/model not found; skipped"
            else:
                retry_count = self.open_part_results(oem)
                if retry_count is None:
                    body_text = normalize(self.page.locator("body").inner_text())
                    if NO_SOURCE_PARTS_RE.search(body_text):
                        result["Status"] = "No sourceable parts available; skipped"
                        log.warning("Skipping VIN=%s OEM=%s: CCC added the estimate line but no sourceable parts were available", vin, oem)
                    else:
                        retries = getattr(self, "vendor_search_retry_count", 0)
                        result["Status"] = f"Vendor search failed after {retries} Retry clicks; skipped"
                else:
                    offers, market_data = self.read_vendor_inventory()
                    result = self.summarize(vin, oem, offers, self.description_target)
                    result["market_data"] = json.dumps(market_data, ensure_ascii=False)
                    log.info("Completed VIN=%s OEM=%s; iPart listed=%s", vin, oem, result["Ipart listed on CCC"])
        except CatalogPartNotFoundError as error:
            log.warning("Skipping VIN=%s OEM=%s: %s", vin, oem, error)
            result["Status"] = f"{error}; skipped"
        except FatalCrawlerError:
            raise
        except Exception as error:
            error_text = str(error).lower()
            browser_closed = self.page.is_closed() or (
                isinstance(error, PlaywrightError) and "closed" in error_text
            )
            overlay_blocked = isinstance(error, PlaywrightError) and "intercepts pointer events" in error_text
            if browser_closed:
                try:
                    self.capture_failure(f"browser_closed_{record['_row']}_{vin}")
                except PlaywrightError:
                    pass
                raise FatalCrawlerError(
                    f"CCC browser closed while processing VIN={vin}, OEM={oem}; stopping crawl to preserve time"
                ) from error
            if overlay_blocked:
                try:
                    self.capture_failure(f"overlay_blocked_{record['_row']}_{vin}")
                except PlaywrightError:
                    pass
                log.warning(
                    "Skipping VIN=%s OEM=%s: CCC overlay blocked interaction; continuing with a fresh workfile",
                    vin,
                    oem,
                )
                result["Status"] = "ERROR: CCC overlay blocked interaction; retry on resume"
                self._workfile_initialized = False
            else:
                log.exception("Row failed: VIN=%s OEM=%s: %s", vin, oem, error)
                self.capture_failure(f"row_{record['_row']}_{vin}")
                result["Status"] = f"ERROR: {error}"
                self._workfile_initialized = False
        finally:
            result["date_time"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        return result


def main() -> int:
    load_dotenv()
    parser = argparse.ArgumentParser(description="Crawl CCC vendor prices for VIN/OEM rows in an Excel file.")
    parser.add_argument("--input", type=Path, default=Path(os.getenv("CCC_INPUT", "input.xlsx")))
    parser.add_argument("--output", type=Path, default=Path(os.getenv("CCC_OUTPUT", "ccc-results.xlsx")))
    parser.add_argument(
        "--description",
        default=os.getenv("CCC_DESCRIPTION", TARGET_DESCRIPTION),
        help=f"Description substring to match when ranking offers (default: {TARGET_DESCRIPTION}).",
    )
    parser.add_argument("--headed", action="store_true", help="Show the browser window (recommended for first run).")
    parser.add_argument("--resume", action="store_true", help="Continue from completed rows in the existing output workbook.")
    args = parser.parse_args()

    username = os.getenv("CCC_USERNAME")
    password = os.getenv("CCC_PASSWORD")
    if not username or not password:
        log.error("Set CCC_USERNAME and CCC_PASSWORD in your environment or .env file.")
        return 2
    if not args.input.is_file():
        log.error("Input workbook not found: %s", args.input)
        return 2

    try:
        records = read_input(args.input)
    except Exception as error:
        log.error("Could not read input workbook: %s", error)
        return 2
    if not records:
        log.error("No rows with both VIN and OEM were found in %s", args.input)
        return 2

    results = load_results(args.output) if args.resume else []
    if args.resume:
        log.info("Resume enabled: loaded %d saved result row(s) from %s", len(results), args.output)
    completed, retry_rows = resume_state(results)
    if results:
        log.info("Loaded %d saved result(s) from %s", len(results), args.output)
    saved_remaining = completed.copy()
    retry_remaining = Counter({key: len(indexes) for key, indexes in retry_rows.items()})
    unfinished = 0
    for record in records:
        key = (record["VIN"], record["OEM"])
        if saved_remaining[key] > 0:
            saved_remaining[key] -= 1
        elif retry_remaining[key] > 0:
            retry_remaining[key] -= 1
        else:
            unfinished += 1
    if unfinished == 0:
        log.info("All %d input rows already have saved results", len(records))
        return 0

    AUTH_STATE.parent.mkdir(parents=True, exist_ok=True)
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(headless=not args.headed)
        context_options: dict[str, Any] = {}
        if AUTH_STATE.is_file():
            context_options["storage_state"] = str(AUTH_STATE)
        context: BrowserContext = browser.new_context(**context_options)
        page = context.new_page()
        page.route("**/*", block_nonessential_resources)
        crawler = CCCCrawler(page, username, password, Path("artifacts"), args.description, context)
        rows_since_save = 0
        try:
            crawler.sign_in_if_needed()
            context.storage_state(path=str(AUTH_STATE))
            crawler.open_workfiles()
            for index, record in enumerate(records, start=1):
                log.info("Progress: %d/%d", index, len(records))
                key = (record["VIN"], record["OEM"])
                if completed[key] > 0:
                    completed[key] -= 1
                    log.info("Skipping row %d; result is already saved", index)
                    continue
                result = crawler.crawl_row(record)
                if retry_rows[key]:
                    results[retry_rows[key].popleft()] = result
                else:
                    results.append(result)
                rows_since_save += 1
                if rows_since_save >= RESULT_SAVE_INTERVAL:
                    save_results(args.output, results)
                    rows_since_save = 0
                    log.info("Saved progress to %s", args.output)
            if rows_since_save:
                save_results(args.output, results)
                log.info("Saved final progress to %s", args.output)
        except Exception as error:
            log.exception("Crawler stopped before completing all rows: %s", error)
            try:
                save_results(args.output, results)
            except Exception:
                log.exception("Could not save partial results to %s", args.output)
            try:
                crawler.capture_failure("crawler_stopped")
            except Exception:
                log.exception("Could not capture failure evidence")
            return 1
        finally:
            context.close()
            browser.close()
    log.info("Crawl complete: %d/%d rows written to %s", len(results), len(records), args.output)
    return 0


if __name__ == "__main__":
    sys.exit(main())